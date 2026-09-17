# -*- coding: utf-8 -*-
import json, collections, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
d=json.load(open("as_1696.json")); J={x["ss"]:x for x in json.load(open("cr_tm_por_ss.json"))}
TMA={x["ss"]:x for x in json.load(open("agosto_tmae.json"))}
CONF={x["ss"]:x for x in json.load(open("confronto_origem.json"))}
base1671=set(CONF)
def U(r,c): return str(r.get(c) or "").strip().upper()
def limpo(v):
    s=str(v or "").strip().upper(); return "" if s in ("","0","00","000","00000","N/A","NA","-","NONE") else s
MES={"2026-01":"jan","2026-02":"fev","2026-03":"mar","2026-04":"abr","2026-05":"mai","2026-06":"jun","2026-07":"jul","2026-08":"ago"}
GAT={"FURTO":"furto","REMANEJAMENTO":"remanejamento","SEM TROCA (NÃO SUBSTITUÍDO)":"sem_troca",
 "DIVISÃO DE CIRCUITO":"divisao","MELHORIA DE POSTO":"melhoria","INCONCLUSIVO":"inconclusivo",
 "SEM OS E SEM OBRA":"sem_os","PREVENTIVO/PROGRAMADO":"preventivo","ABALROAMENTO":"abalroamento",
 "TAPE/REGULARIZAÇÃO DE TENSÃO":"tap","AUXILIAR DE RELIGADOR":"auxiliar","AUSENTE DA CRÍTICA":"sem_interrupcao",
 "FORA DA JANELA DA CRÍTICA":"fora_da_janela","ERRO DE CADASTRO — NÃO É TRANSFORMADOR":"erro_cadastro",
 "DANO DE TERCEIROS":"terceiros","FALTA DE FASE":"falta_fase","SS DUPLICADA":"duplicada",
 "OBRA SEM TRANSFORMADOR NO MATERIAL":"obra_sem_transformador","OBRA SEM EXECUÇÃO":"obra_sem_execucao",
 "AVALIAR COM O MATHEUS":"avaliar_matheus","SEM OBRA — PASSADOS 60 DIAS":"sem_obra",
 "RETIDO — SEM PROVA DE TROCA":"sem_prova_de_troca","RETIDO — SEM INTERRUPÇÃO NA JANELA":"sem_interrupcao",
 "RETIDO — RESSALVA DA INTERRUPÇÃO":"ressalva_da_interrupcao"}
linhas=[]; novos=[]
for r in d:
    ss=r["ss"]
    if ss not in base1671: continue
    m=MES.get(str(r.get("mes")),"")
    cat=U(r,"categoria"); conta=U(r,"conta")=="SIM"
    j=J.get(ss,{}); t=TMA.get(ss,{})
    troca = bool(limpo(r.get("ns_ret")) and limpo(r.get("ns_inst")) and limpo(r.get("ns_ret"))!=limpo(r.get("ns_inst")))
    origem="leitura caso a caso"
    if conta:
        if m=="ago" and not j.get("cr_causa"):
            if troca: res="RETIDO"; gat="sem_interrupcao"; cat="RETIDO — SEM INTERRUPÇÃO NA JANELA"
            else: res="EXPURGO"; gat="sem_interrupcao"; cat="AUSENTE DA CRÍTICA"
            origem="Crítica + TMAE de agosto (novo)"
            novos.append(ss)
        else:
            continue
    else:
        res="RETIDO" if cat.startswith("RETIDO") else "EXPURGO"
        gat=GAT.get(cat,"outro")
    linhas.append({"ss":ss,"mes":m,"resultado":res,"gatilho":gat,"categoria":cat,
      "trafo":r.get("trafo"),"os":r.get("os"),"obra":r.get("obra"),
      "abertura":str(r.get("abertura"))[:16],"termino":str(r.get("termino"))[:16],
      "tipo":r.get("tipo_ss"),"equipe":r.get("equipe"),"localidade":r.get("localidade"),
      "origem_ss":r.get("origem_ss"),"defeito_ss":r.get("defeito_ss"),
      "ns_ret":r.get("ns_ret"),"ns_inst":r.get("ns_inst"),"prova_troca":"sim" if troca else "não",
      "cr_causa":j.get("cr_causa",""),"cr_sub":j.get("cr_sub",""),"cr_papel":j.get("cr_papel",""),
      "tm_causa":j.get("tm_causa","") or t.get("causa",""),"tm_sub":j.get("tm_sub","") or t.get("sub",""),
      "no_infotrafo":r.get("no_infotrafo"),"origem_leitura":origem,"motivo":r.get("motivo",""),
      "tss":(CONF[ss]["tss"] or "")[:400],"tos":(CONF[ss]["tos"] or "")[:600]})
print("total fora do indicador:",len(linhas),"| novos de agosto:",len(novos))
c=collections.Counter(x["resultado"] for x in linhas); print(c.most_common())
print("por mês:",sorted(collections.Counter(x["mes"] for x in linhas).items(),key=lambda x:["jan","fev","mar","abr","mai","jun","jul","ago"].index(x[0])))
print("gatilhos:",collections.Counter(x["gatilho"] for x in linhas).most_common(25))
NAVY="1F3864"; borda=Border(*[Side(style="thin",color="D9D9D9")]*4)
wb=openpyxl.Workbook()
ordemmes={m:i for i,m in enumerate(["jan","fev","mar","abr","mai","jun","jul","ago"])}
def aba(nome,dados):
    ws=wb.create_sheet(nome)
    cols=[("SS","ss"),("Mês","mes"),("Resultado","resultado"),("Gatilho","gatilho"),("Categoria","categoria"),
     ("Trafo","trafo"),("OS","os"),("Obra","obra"),("Abertura","abertura"),("Término","termino"),
     ("Tipo da SS","tipo"),("Equipe","equipe"),("Localidade","localidade"),
     ("Origem da SS","origem_ss"),("Defeito da SS","defeito_ss"),
     ("NS retirado","ns_ret"),("NS instalado","ns_inst"),("Prova de troca","prova_troca"),
     ("Crítica — causa","cr_causa"),("Crítica — subcausa","cr_sub"),("Crítica — papel","cr_papel"),
     ("TMAE — causa","tm_causa"),("TMAE — subcausa","tm_sub"),
     ("No Infotrafo","no_infotrafo"),("Origem da leitura","origem_leitura"),("Motivo","motivo"),
     ("Texto da SS","tss"),("Texto da OS","tos")]
    ws.append([a for a,_ in cols])
    for c_ in ws[1]:
        c_.font=Font(bold=True,color="FFFFFF",size=10); c_.fill=PatternFill("solid",fgColor=NAVY)
        c_.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for x in sorted(dados,key=lambda x:(ordemmes[x["mes"]],x["gatilho"],x["ss"])):
        ws.append([x[k] for _,k in cols])
        if x["origem_leitura"].startswith("Crítica"):
            for i in range(1,len(cols)+1): ws.cell(row=ws.max_row,column=i).fill=PatternFill("solid",fgColor="FFF2CC")
    for i,(a,_) in enumerate(cols,1):
        ws.column_dimensions[get_column_letter(i)].width={"Texto da SS":50,"Texto da OS":65,"Motivo":45,
          "Categoria":30,"Crítica — subcausa":30,"TMAE — subcausa":28,"Origem da leitura":26,"SS":22,"Tipo da SS":26}.get(a,13)
    for row in ws.iter_rows(min_row=2):
        for c_ in row: c_.border=borda; c_.alignment=Alignment(vertical="top",wrap_text=c_.column in (26,27,28))
    ws.freeze_panes="B2"; ws.auto_filter.ref=ws.dimensions; ws.row_dimensions[1].height=36
    return ws
aba("Expurgos",[x for x in linhas if x["resultado"]=="EXPURGO"])
aba("Retidos",[x for x in linhas if x["resultado"]=="RETIDO"])
w=wb.create_sheet("Resumo",0)
w.append(["Base de expurgos — janeiro a agosto de 2026",""])
w.append(["",""])
w.append(["Universo (SS de substituição de transformador)",1671])
w.append(["Contam no indicador",1671-len(linhas)])
w.append(["Fora do indicador",len(linhas)])
w.append(["   Expurgos",c["EXPURGO"]])
w.append(["   Retidos (podem voltar)",c["RETIDO"]])
w.append(["",""])
w.append(["Novos nesta atualização (Crítica + TMAE de agosto)",len(novos)])
w.append(["",""])
w.append(["Expurgos por mês",""])
for m,v in sorted(collections.Counter(x["mes"] for x in linhas if x["resultado"]=="EXPURGO").items(),key=lambda x:ordemmes[x[0]]):
    w.append(["   "+m,v])
w.append(["",""])
w.append(["Expurgos por gatilho",""])
for k,v in collections.Counter(x["gatilho"] for x in linhas if x["resultado"]=="EXPURGO").most_common():
    w.append(["   "+k,v])
w.append(["",""])
w.append(["Retidos por gatilho",""])
for k,v in collections.Counter(x["gatilho"] for x in linhas if x["resultado"]=="RETIDO").most_common():
    w.append(["   "+k,v])
w["A1"].font=Font(bold=True,size=13,color=NAVY)
for r_ in (3,4,5,6,7,9,11,len(w["A"])): pass
for row in w.iter_rows(min_row=1):
    if row[0].value and not str(row[0].value).startswith("   "): row[0].font=Font(bold=True,size=11 if row[0].row>1 else 13,color=NAVY if row[0].row==1 else "000000")
w.column_dimensions["A"].width=52; w.column_dimensions["B"].width=12
del wb["Sheet"]
wb.save("Base_de_Expurgos_jan_ago.xlsx")
json.dump(linhas,open("base_expurgos.json","w"),ensure_ascii=False)
print("ok")
