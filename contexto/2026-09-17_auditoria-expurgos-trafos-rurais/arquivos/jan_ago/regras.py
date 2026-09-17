# -*- coding: utf-8 -*-
import json, collections
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

BASE = "/tmp/claude-0/-home-user/74dc9c64-5026-54ee-a81e-173d2f38a735/scratchpad/jan_ago/"
dados = json.load(open(BASE + "as_1696.json"))
cont = collections.Counter((r.get("categoria") or "").strip().upper() for r in dados)

from textos import REGRAS, COMO_LER

AZUL = "1F3864"; CINZA = "D9E2F3"; VERDE = "E2EFDA"; VERM = "FCE4E4"; AMAR = "FFF2CC"
borda = Border(*[Side(style="thin", color="BFBFBF")] * 4)

wb = Workbook()
ws = wb.active
ws.title = "Regras"
ws.append(["Categoria", "Descrição", "SS jan-ago/2026"])
for i, c in enumerate(ws[1], 1):
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = PatternFill("solid", fgColor=AZUL)
    c.alignment = Alignment(horizontal="center", vertical="center")

def norm(s):
    import unicodedata
    return unicodedata.normalize("NFD", s).encode("ascii", "ignore").decode().upper()

chave = {norm(k): v for k, v in cont.items()}

linha = 1
for cat, desc in REGRAS:
    linha += 1
    if desc is None:
        ws.append([cat, "", ""])
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=3)
        cel = ws.cell(row=linha, column=1)
        cel.font = Font(bold=True, color="FFFFFF", size=10)
        cel.fill = PatternFill("solid", fgColor="4472C4")
        cel.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[linha].height = 22
        continue
    n = chave.get(norm(cat), 0)
    ws.append([cat, desc, n])
    cor = VERDE if "CONTAM NO INDICADOR" in "" else None
    ws.cell(row=linha, column=1).font = Font(bold=True, size=10)
    ws.cell(row=linha, column=1).alignment = Alignment(vertical="top", wrap_text=True)
    ws.cell(row=linha, column=2).alignment = Alignment(vertical="top", wrap_text=True)
    ws.cell(row=linha, column=3).alignment = Alignment(horizontal="center", vertical="top")
    for col in range(1, 4):
        ws.cell(row=linha, column=col).border = borda
    ws.row_dimensions[linha].height = 45

# pinta o bloco que conta de verde
secao = None
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value or ""
    if ws.cell(row=r, column=2).value in (None, ""):
        secao = v
        continue
    if secao and secao.startswith("CONTAM"):
        cor = VERDE
    elif secao and secao.startswith("AINDA"):
        cor = AMAR
    else:
        cor = VERM
    for col in range(1, 4):
        ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=cor)

ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 105
ws.column_dimensions["C"].width = 16
ws.freeze_panes = "A2"

# aba Como ler
w2 = wb.create_sheet("Como ler")
w2.append(["Princípio", "Explicação"])
for c in w2[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=AZUL)
for a, b in COMO_LER:
    w2.append([a, b])
for r in range(2, w2.max_row + 1):
    w2.cell(row=r, column=1).font = Font(bold=True, size=10)
    w2.cell(row=r, column=1).alignment = Alignment(vertical="top", wrap_text=True)
    w2.cell(row=r, column=2).alignment = Alignment(vertical="top", wrap_text=True)
    w2.row_dimensions[r].height = 42
w2.column_dimensions["A"].width = 40
w2.column_dimensions["B"].width = 105

# aba Resumo
w3 = wb.create_sheet("Resumo")
conta_sim = sum(v for k, v in cont.items() if k in ("QUEIMADO", "AVARIADO"))
w3.append(["Bloco", "SS"])
for c in w3[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=AZUL)
w3.append(["Universo jan-ago/2026 (1.444 Infotrafo + 252 fora)", len(dados)])
w3.append(["Contam no indicador (queimado + avariado)", conta_sim])
w3.append(["Nao contam", len(dados) - conta_sim])
for r in range(2, w3.max_row + 1):
    w3.cell(row=r, column=1).font = Font(bold=(r == 2))
w3.column_dimensions["A"].width = 55
w3.column_dimensions["B"].width = 12

wb.save(BASE + "Regras_de_Categorizacao.xlsx")
faltando = [k for k in cont if norm(k) not in {norm(a) for a, b in REGRAS if b}]
print("ok", len(dados), conta_sim, "sem regra:", faltando)
