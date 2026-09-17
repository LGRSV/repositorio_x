# -*- coding: utf-8 -*-
import json, collections, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from novas23 import NOVAS

rows=json.load(open("base_1671.json"))
d=json.load(open("as_1696.json"))
cov={r["ss"]:r for r in d}
def S(r,c): return str(r.get(c) or "").strip()
MES={"01":"jan","02":"fev","03":"mar","04":"abr","05":"mai","06":"jun","07":"jul","08":"ago","09":"set"}
def mes(v):
    v=str(v or "")
    return MES.get(v[5:7],"") if v[4:5]=="-" else MES.get(v[3:5],"")
def limpo(v):
    s=str(v or "").strip().upper()
    return "" if s in ("","0","00","000","00000","N/A","NA","NONE","-") else s

out=[]
for r in rows:
    ss=S(r,"NUMERO_SS")
    if ss in cov:
        a=cov[ss]
        cat=str(a.get("categoria") or "").strip()
        conta="SIM" if str(a.get("conta") or "").strip().upper()=="SIM" else "NÃO"
        fonte="site / análise já feita"
        motivo=str(a.get("motivo") or "")
        evid=""; conf=str(a.get("confianca") or "")
    elif ss in NOVAS:
        cat,conta,motivo,evid,conf=NOVAS[ss]
        fonte="leitura nova (SS + OS)"
    else:
        cat,conta,motivo,evid,conf,fonte="SEM LEITURA","NÃO","","","","—"
    out.append(collections.OrderedDict([
        ("NUMERO_SS",ss),("NUMERO_OS",S(r,"NUMERO_OS")),("NUM_OBRA",S(r,"NUM_OBRA")),("NUM_TRAFO",S(r,"NUM_TRAFO")),
        ("TIPOSS",S(r,"TIPOSS")),("Mês",mes(S(r,"DATA_ABERTURA_SS"))),
        ("DATA_ABERTURA_SS",S(r,"DATA_ABERTURA_SS")[:16]),("DATA_TERMINO_SS",S(r,"DATA_TERMINO_SS")[:16]),
        ("SITUACAO_SS",S(r,"SITUACAO_SS")),("CRITICIDADE_SS",S(r,"CRITICIDADE_SS")),
        ("ORIGEM_SS",S(r,"ORIGEM_SS")),("DEFEITO_SS",S(r,"DEFEITO_SS")),
        ("ORIGEM (dinâmica atual)",S(r,"ORIGEM")),("DEFEITO (dinâmica atual)",S(r,"DEFEITO")),
        ("CATEGORIA RECATEGORIZADA",cat),("Conta no indicador",conta),
        ("Fonte da categoria",fonte),("Confiança",conf),
        ("Motivo",motivo),("Evidência",evid),
        ("NS_RETIRADO",S(r,"NS_RETIRADO")),("NS_INSTALADO",S(r,"NS_INSTALADO")),
        ("POTENCIA_RET",S(r,"POTENCIA_RET")),("POTENCIA_INST",S(r,"POTENCIA_INST")),
        ("PROVÁVEL_MOTIVO_DO_DEFEITO",S(r,"PROVÁVEL_MOTIVO_DO_DEFEITO")),
        ("COD_EQUIPE",S(r,"COD_EQUIPE")),("LOCALIDADE",S(r,"LOCALIDADE")),("RURA_URBANO",S(r,"RURA_URBANO")),
        ("DESCRIPTION_SS",S(r,"DESCRIPTION_SS")[:600].replace("_x000D_"," ")),
        ("DESCRICAO_OS",S(r,"DESCRICAO_OS")[:900].replace("_x000D_"," ")),
    ]))
_TODAS=list(out)
out=[x for x in out if x["Mês"]!="set"]
print("linhas",len(out))
print(collections.Counter(x["Fonte da categoria"] for x in out))

NAVY="1F3864"; AZUL="4472C4"
borda=Border(*[Side(style="thin",color="D9D9D9")]*4)
wb=openpyxl.Workbook()

# ---------- aba base
ws=wb.active; ws.title="Base recategorizada"
cols=list(out[0].keys()); ws.append(cols)
for c in ws[1]:
    c.font=Font(bold=True,color="FFFFFF",size=10); c.fill=PatternFill("solid",fgColor=NAVY)
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
for x in out: ws.append([x[c] for c in cols])
for i,c in enumerate(cols,1):
    L=get_column_letter(i)
    ws.column_dimensions[L].width = 60 if c in ("DESCRIPTION_SS","DESCRICAO_OS") else (34 if c in ("Motivo","Evidência") else max(11,min(22,len(c)+3)))
ws.freeze_panes="B2"; ws.auto_filter.ref=ws.dimensions; ws.row_dimensions[1].height=36
ic=cols.index("CATEGORIA RECATEGORIZADA")+1
CONTAM={"QUEIMADO","AVARIADO"}
for r in range(2,ws.max_row+1):
    v=str(ws.cell(row=r,column=ic).value or "")
    cor="E2EFDA" if v in CONTAM else ("FFF2CC" if v.startswith(("RETIDO","PENDENTE","INCONCLUSIVO","AVALIAR")) else "FCE4E4")
    ws.cell(row=r,column=ic).fill=PatternFill("solid",fgColor=cor)
    ws.cell(row=r,column=ic).font=Font(bold=True,size=10)

# ---------- dinâmica nova: TIPOSS x categoria x mês
meses=["jan","fev","mar","abr","mai","jun","jul","ago"]
w2=wb.create_sheet("Dinamica nova")
w2.append(["TIPOSS","Categoria recategorizada","Conta"]+meses+["Total"])
for c in w2[1]:
    c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY)
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
piv=collections.defaultdict(lambda: collections.Counter())
for x in out: piv[(x["TIPOSS"],x["CATEGORIA RECATEGORIZADA"],x["Conta no indicador"])][x["Mês"]]+=1
def ordem(k):
    t,c,_=k
    return (0 if t.startswith("FORMS") else (1 if t.startswith("AVISO") else 2), -sum(piv[k].values()))
for k in sorted(piv,key=ordem):
    t,c,co=k; cnt=piv[k]
    w2.append([t,c,co]+[cnt.get(m,"") for m in meses]+[sum(cnt.values())])
    if co=="SIM":
        for i in range(1,4+len(meses)+1): w2.cell(row=w2.max_row,column=i).fill=PatternFill("solid",fgColor="E2EFDA")
tot=collections.Counter()
for x in out: tot[x["Mês"]]+=1
w2.append(["TOTAL GERAL","","",]+[tot.get(m,"") for m in meses]+[len(out)])
for c in w2[w2.max_row]: c.font=Font(bold=True)
for L,w in zip(["A","B","C"],(32,38,8)): w2.column_dimensions[L].width=w
for i in range(4,4+len(meses)+1): w2.column_dimensions[get_column_letter(i)].width=7
w2.freeze_panes="D2"; w2.row_dimensions[1].height=30

# ---------- de -> para
w3=wb.create_sheet("De para")
cats=sorted({x["CATEGORIA RECATEGORIZADA"] for x in out})
w3.append(["ORIGEM na dinâmica atual"]+cats+["Total"])
for c in w3[1]:
    c.font=Font(bold=True,color="FFFFFF",size=9); c.fill=PatternFill("solid",fgColor=NAVY)
    c.alignment=Alignment(horizontal="center",vertical="bottom",wrap_text=True,text_rotation=60)
w3.cell(row=1,column=1).alignment=Alignment(horizontal="left",vertical="bottom",wrap_text=True)
m=collections.defaultdict(collections.Counter)
for x in out: m[x["ORIGEM (dinâmica atual)"] or "(vazio)"][x["CATEGORIA RECATEGORIZADA"]]+=1
for o in sorted(m,key=lambda o:-sum(m[o].values())):
    w3.append([o]+[m[o].get(c,"") for c in cats]+[sum(m[o].values())])
w3.column_dimensions["A"].width=30
for i in range(2,len(cats)+3): w3.column_dimensions[get_column_letter(i)].width=6
w3.row_dimensions[1].height=150; w3.freeze_panes="B2"

# ---------- novas
w4=wb.create_sheet("As 23 de setembro (fora)")
import json as _j
sub=[]
for _x in _TODAS:
    if _x["Fonte da categoria"].startswith("leitura nova"): sub.append(_x)
c4=["NUMERO_SS","TIPOSS","DATA_ABERTURA_SS","SITUACAO_SS","ORIGEM (dinâmica atual)","DEFEITO (dinâmica atual)","CATEGORIA RECATEGORIZADA","Conta no indicador","Confiança","Motivo","Evidência","NS_RETIRADO","NS_INSTALADO","DESCRIPTION_SS","DESCRICAO_OS"]
w4.append(c4)
for c in w4[1]:
    c.font=Font(bold=True,color="FFFFFF",size=10); c.fill=PatternFill("solid",fgColor=NAVY)
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
for x in sub: w4.append([x[c] for c in c4])
for i,c in enumerate(c4,1):
    L=get_column_letter(i)
    w4.column_dimensions[L].width = 60 if c in ("DESCRIPTION_SS","DESCRICAO_OS") else (44 if c in ("Motivo","Evidência") else max(11,min(22,len(c)+3)))
for r in w4.iter_rows(min_row=2):
    for c in r: c.alignment=Alignment(vertical="top",wrap_text=True); c.border=borda
w4.freeze_panes="B2"; w4.row_dimensions[1].height=32

# ---------- resumo
w5=wb.create_sheet("Resumo",0)
w5.append(["Categoria recategorizada","SS","Conta no indicador"])
for c in w5[1]:
    c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY)
cc=collections.Counter(x["CATEGORIA RECATEGORIZADA"] for x in out)
conta_map={x["CATEGORIA RECATEGORIZADA"]:x["Conta no indicador"] for x in out}
for k,v in cc.most_common():
    w5.append([k,v,conta_map.get(k,"")])
    if conta_map.get(k)=="SIM":
        for i in range(1,4): w5.cell(row=w5.max_row,column=i).fill=PatternFill("solid",fgColor="E2EFDA")
w5.append([]); w5.append(["TOTAL",len(out),""])
w5.append(["Contam (queimado + avariado)",sum(1 for x in out if x["Conta no indicador"]=="SIM"),""])
w5.append(["Não contam",sum(1 for x in out if x["Conta no indicador"]!="SIM"),""])
for r in range(w5.max_row-2,w5.max_row+1):
    w5.cell(row=r,column=1).font=Font(bold=True)
w5.column_dimensions["A"].width=42; w5.column_dimensions["B"].width=10; w5.column_dimensions["C"].width=20
wb.save("Dinamica_defeito_jan_ago.xlsx")
print("contam:",sum(1 for x in out if x["Conta no indicador"]=="SIM"))
for k,v in cc.most_common(10): print("  %4d %s"%(v,k))
