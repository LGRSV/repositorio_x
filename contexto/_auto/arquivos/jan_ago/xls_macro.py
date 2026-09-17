# -*- coding: utf-8 -*-
import json, collections, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
L=json.load(open("base_expurgos.json"))
NAVY="1F3864"; borda=Border(*[Side(style="thin",color="D9D9D9")]*4)
ORD=["Sem interrupção comprovada","Causa externa ao transformador","Troca sem falha","Em aberto","Sem documento","Não houve troca"]
CORM={"Sem interrupção comprovada":"E8EDF5","Causa externa ao transformador":"FCE4E4","Troca sem falha":"FDEBD3",
      "Em aberto":"FFF2CC","Sem documento":"EDE7F6","Não houve troca":"E2EFDA"}
MESO=["jan","fev","mar","abr","mai","jun","jul","ago"]
wb=openpyxl.Workbook()
def aba(nome,dados):
    ws=wb.create_sheet(nome)
    cols=[("SS","ss"),("Mês","mes"),("Macro categoria","macro"),("Motivo","rotulo"),("Resultado","resultado"),
     ("Categoria pelo texto","cat_texto"),("Fonte da categoria","cat_texto_fonte"),
     ("Categoria da leitura","categoria"),("Trafo","trafo"),("OS","os"),("Obra","obra"),
     ("Abertura","abertura"),("Término","termino"),("Tipo da SS","tipo"),("Equipe","equipe"),("Localidade","localidade"),
     ("Origem da SS","origem_ss"),("Defeito da SS","defeito_ss"),("NS retirado","ns_ret"),("NS instalado","ns_inst"),
     ("Prova de troca","prova_troca"),("Crítica — causa","cr_causa"),("Crítica — subcausa","cr_sub"),
     ("TMAE — causa","tm_causa"),("TMAE — subcausa","tm_sub"),("No Infotrafo","no_infotrafo"),
     ("Origem da leitura","origem_leitura"),("Por que saiu","motivo"),("Texto da SS","tss"),("Texto da OS","tos")]
    ws.append([a for a,_ in cols])
    for c_ in ws[1]:
        c_.font=Font(bold=True,color="FFFFFF",size=10); c_.fill=PatternFill("solid",fgColor=NAVY)
        c_.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for x in sorted(dados,key=lambda x:(ORD.index(x["macro"]),x["rotulo"],MESO.index(x["mes"]),x["ss"])):
        ws.append([x.get(k) for _,k in cols])
        ws.cell(row=ws.max_row,column=3).fill=PatternFill("solid",fgColor=CORM[x["macro"]])
        if str(x.get("origem_leitura","")).startswith("Crítica"):
            ws.cell(row=ws.max_row,column=25).fill=PatternFill("solid",fgColor="FFF2CC")
    for i,(a,_) in enumerate(cols,1):
        ws.column_dimensions[get_column_letter(i)].width={"Texto da SS":50,"Texto da OS":65,"Por que saiu":45,
          "Macro categoria":30,"Categoria da leitura":30,"Categoria pelo texto":22,"Fonte da categoria":26,"Crítica — subcausa":30,"TMAE — subcausa":28,
          "Origem da leitura":26,"SS":22,"Tipo da SS":26,"Motivo":22}.get(a,13)
    for row in ws.iter_rows(min_row=2):
        for c_ in row: c_.border=borda; c_.alignment=Alignment(vertical="top",wrap_text=c_.column in (26,27,28))
    ws.freeze_panes="B2"; ws.auto_filter.ref=ws.dimensions; ws.row_dimensions[1].height=36
w=wb.create_sheet("Resumo")
w.append(["Base de expurgos — janeiro a agosto de 2026"])
w["A1"].font=Font(bold=True,size=14,color=NAVY)
w.append([])
w.append(["Macro categoria","Expurgo","Retido","Total"])
for c_ in w[3]:
    c_.font=Font(bold=True,color="FFFFFF"); c_.fill=PatternFill("solid",fgColor=NAVY); c_.alignment=Alignment(horizontal="center")
m=collections.defaultdict(collections.Counter)
for x in L: m[x["macro"]][x["resultado"]]+=1
for k in ORD:
    w.append([k,m[k]["EXPURGO"],m[k]["RETIDO"],sum(m[k].values())])
    w.cell(row=w.max_row,column=1).fill=PatternFill("solid",fgColor=CORM[k])
tot=collections.Counter(x["resultado"] for x in L)
w.append(["TOTAL",tot["EXPURGO"],tot["RETIDO"],len(L)])
for c_ in w[w.max_row]: c_.font=Font(bold=True)
w.append([]); w.append(["Detalhe por motivo","Expurgo","Retido","Total"])
for c_ in w[w.max_row]:
    c_.font=Font(bold=True,color="FFFFFF"); c_.fill=PatternFill("solid",fgColor="4472C4"); c_.alignment=Alignment(horizontal="center")
for k in ORD:
    for g in sorted({x["rotulo"] for x in L if x["macro"]==k},key=lambda g:-sum(1 for x in L if x["rotulo"]==g)):
        e=sum(1 for x in L if x["rotulo"]==g and x["resultado"]=="EXPURGO")
        r=sum(1 for x in L if x["rotulo"]==g and x["resultado"]=="RETIDO")
        w.append(["   "+g,e,r,e+r])
        w.cell(row=w.max_row,column=1).fill=PatternFill("solid",fgColor=CORM[k])
w.append([]); w.append(["Expurgos por mês",""])
w.cell(row=w.max_row,column=1).font=Font(bold=True)
for mm in MESO:
    w.append(["   "+mm,sum(1 for x in L if x["mes"]==mm and x["resultado"]=="EXPURGO"),
              sum(1 for x in L if x["mes"]==mm and x["resultado"]=="RETIDO"),
              sum(1 for x in L if x["mes"]==mm)])
w.append([]); w.append(["O que o equipamento era, pelo texto","Expurgo","Retido","Total"])
for c_ in w[w.max_row]:
    c_.font=Font(bold=True,color="FFFFFF"); c_.fill=PatternFill("solid",fgColor="4472C4"); c_.alignment=Alignment(horizontal="center")
for k,v in collections.Counter(x["cat_texto"] for x in L).most_common():
    w.append(["   "+k,sum(1 for x in L if x["cat_texto"]==k and x["resultado"]=="EXPURGO"),
              sum(1 for x in L if x["cat_texto"]==k and x["resultado"]=="RETIDO"),v])
w.append([]); w.append(["Universo",1671]); w.append(["Contam no indicador",1671-len(L)])
w.append(["Novos nesta atualização (Crítica + TMAE de agosto)",sum(1 for x in L if str(x.get("origem_leitura","")).startswith("Crítica"))])
for r_ in w.iter_rows(min_row=1):
    if r_[0].value and not str(r_[0].value).startswith("   ") and r_[0].row>3: r_[0].font=Font(bold=True)
for L_,wd in zip("ABCD",(46,11,11,11)): w.column_dimensions[L_].width=wd
aba("Expurgos",[x for x in L if x["resultado"]=="EXPURGO"])
aba("Retidos",[x for x in L if x["resultado"]=="RETIDO"])
del wb["Sheet"]
wb.move_sheet("Resumo",offset=-2)
wb.save("Base_de_Expurgos_jan_ago.xlsx"); print("ok",tot["EXPURGO"],tot["RETIDO"])
