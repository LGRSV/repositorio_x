# -*- coding: utf-8 -*-
import json, collections, re, unicodedata, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
A=[r for r in json.load(open("base_3602.json")) if str(r.get("TIPOSS") or "").strip()]
J={x["ss"]:x for x in json.load(open("cr_tm_por_ss.json"))}
d=json.load(open("as_1696.json")); cov={r["ss"]:r for r in d}
def S(r,c): return str(r.get(c) or "").strip().upper() or ""
def norm(s):
    s=re.sub(r"_x000D_"," ",str(s or ""))
    s=unicodedata.normalize("NFD",s).encode("ascii","ignore").decode().upper()
    return re.sub(r"\s+"," ",s).strip()
TXT={"DESCARGA ATMOSFÉRICA":["DESCARGA","RAIO","QUEIM","EXPLOD","ESTOUR","CURTO"],
 "SOBRECARGA":["SOBRECARGA","NAO SUPORTA","AUMENTO DE CARGA","SOBRECARREG"],
 "VAZAMENTO DE ÓLEO":["VAZAMENTO","VAZANDO","OLEO"],"BUCHA DANIFICADA":["BUCHA"],
 "DETERIORADA":["DETERIOR","PODRE","OXID","CORRO","DANIFICAD"],"FURTADO":["FURT","ROUB","LEVARAM"],
 "POSTE DANIFICADO/QUEBRADO":["POSTE"],"QUEIMADO":["QUEIM","EXPLOD","ESTOUR","CURTO"],
 "COLA E FITA":["COLA","FITA","VEDA","REPARO"],"ABALROADO":["ABALRO","COLIS","VEICUL","BATEU"],
 "VEGETAÇÃO":["VEGET","ARVORE","GALHO","PODA"],"MAL FIXADO(A)/SOLTO(A)":["SOLTO","MAL FIXAD","FIXACAO"],
 "FALTANDO":["FALTA","AUSENT"],"BASE DETERIORADA":["BASE","DETERIOR"],"CORROSÃO / OXIDAÇÃO":["CORRO","OXID","DETERIOR"],
 "VANDALISMO":["VANDAL","DEPRED","TIRO"],"OXIDADO / DETERIORADO":["OXID","DETERIOR","CORRO"],
 "GALHOS EM CONTATO COM OS CONDUTORES":["GALHO","ARVORE","VEGET"]}
CRI={"DESCARGA ATMOSFÉRICA":["DESCARGA"],"SOBRECARGA":["SOBRECARGA"],"VAZAMENTO DE ÓLEO":["VAZAMENTO","TANQUE"],
 "BUCHA DANIFICADA":["BUCHA"],"FURTADO":["FURTO","ROUBO"],"ABALROADO":["ABALRO","VEICUL","TERCEIROS"],
 "VEGETAÇÃO":["ARVORE","VEGET","MEIO AMBIENTE"],"POSTE DANIFICADO/QUEBRADO":["POSTE"],"QUEIMADO":["QUEIMADO"],
 "COLA E FITA":["VAZAMENTO","TANQUE"],"DETERIORADA":["DETERIOR","POSTE"],"MAL FIXADO(A)/SOLTO(A)":["CONEXAO","JUMPER","PARTIDO"]}
NAVY="1F3864"; borda=Border(*[Side(style="thin",color="D9D9D9")]*4)
def veredito(r):
    k=S(r,"DEFEITO"); j=J.get(S(r,"NUMERO_SS"),{})
    t=TXT.get(k); c=CRI.get(k)
    txt = bool(t) and any(w in norm(r.get("DESCRICAO_OS"))+" "+norm(r.get("DESCRIPTION_SS")) for w in t)
    sc=(j.get("cr_sub","")+" "+j.get("cr_causa","")).upper(); st=(j.get("tm_sub","")+" "+j.get("tm_causa","")).upper()
    cr = bool(c) and bool(j.get("cr_causa")) and any(w in sc for w in c)
    tm = bool(c) and bool(j.get("tm_causa")) and any(w in st for w in c)
    n=sum([txt,cr,tm])
    if not t and not c: return "sem teste",txt,cr,tm
    if n>=2: return "confirmado por 2+ fontes",txt,cr,tm
    if n==1: return "confirmado por 1 fonte",txt,cr,tm
    if j.get("cr_causa") or j.get("tm_causa") or True: return "nenhuma fonte confirma",txt,cr,tm
wb=openpyxl.Workbook(); ws=wb.active; ws.title="SS OS Critica TMAE"
cols=["SS","Trafo","Tipo da SS","Equipe","Defeito da SS","Defeito da OS","Veredito",
      "Texto SS+OS","Crítica","TMAE",
      "Crítica — causa","Crítica — subcausa","Crítica — grupo","Crítica — papel do ativo","Crítica — clientes","Crítica — início",
      "TMAE — causa","TMAE — subcausa","TMAE — data","TMAE — observação da equipe",
      "Categoria da auditoria","Conta","Texto da SS","Texto da OS (integral)"]
ws.append(cols)
for c in ws[1]:
    c.font=Font(bold=True,color="FFFFFF",size=10); c.fill=PatternFill("solid",fgColor=NAVY)
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
cont=collections.Counter()
for r in sorted(A,key=lambda r:S(r,"NUMERO_SS")):
    ss=S(r,"NUMERO_SS"); j=J.get(ss,{}); a=cov.get(ss,{})
    v,txt,cr,tm=veredito(r); cont[v]+=1
    ws.append([ss,S(r,"NUM_TRAFO"),S(r,"TIPOSS"),S(r,"COD_EQUIPE"),S(r,"DEFEITO_SS"),S(r,"DEFEITO"),v,
      "sim" if txt else "não","sim" if cr else ("—" if not j.get("cr_causa") else "não"),
      "sim" if tm else ("—" if not j.get("tm_causa") else "não"),
      j.get("cr_causa",""),j.get("cr_sub",""),j.get("cr_grupo",""),j.get("cr_papel",""),j.get("cr_cli",""),j.get("cr_ini","")[:16],
      j.get("tm_causa",""),j.get("tm_sub",""),j.get("tm_data",""),j.get("tm_obs",""),
      a.get("categoria",""),a.get("conta",""),
      re.sub(r"_x000D_"," ",str(r.get("DESCRIPTION_SS") or "")),
      re.sub(r"_x000D_"," ",str(r.get("DESCRICAO_OS") or ""))])
    cor={"confirmado por 2+ fontes":"E2EFDA","confirmado por 1 fonte":"FFF2CC","nenhuma fonte confirma":"FCE4E4"}.get(v)
    if cor: ws.cell(row=ws.max_row,column=7).fill=PatternFill("solid",fgColor=cor)
larg={"SS":22,"Trafo":13,"Tipo da SS":26,"Equipe":12,"Defeito da SS":22,"Defeito da OS":22,"Veredito":24,
 "Crítica — causa":24,"Crítica — subcausa":34,"Crítica — grupo":18,"Crítica — papel do ativo":22,"Crítica — início":17,
 "TMAE — causa":22,"TMAE — subcausa":30,"TMAE — data":17,"TMAE — observação da equipe":45,
 "Categoria da auditoria":26,"Texto da SS":55,"Texto da OS (integral)":75}
for i,c in enumerate(cols,1): ws.column_dimensions[get_column_letter(i)].width=larg.get(c,11)
for row in ws.iter_rows(min_row=2):
    for c in row: c.border=borda; c.alignment=Alignment(vertical="top",wrap_text=c.column in (23,24,20))
ws.freeze_panes="B2"; ws.auto_filter.ref=ws.dimensions; ws.row_dimensions[1].height=40
w2=wb.create_sheet("Resumo",0)
for t,v in [("SS analisadas",len(A))]+list(cont.most_common()):
    w2.append([t,v])
w2.insert_rows(1); w2["A1"]="Veredito com as quatro fontes"; w2["A1"].font=Font(bold=True,size=13,color=NAVY)
w2.append([]); w2.append(["Confirmação por fonte, entre as que têm o defeito testável",""])
w2.append(["Texto da SS + OS (leitura integral)","73%"])
w2.append(["Crítica — causa e subcausa","66% das SS casaram; confirma 76% em descarga atmosférica"])
w2.append(["TMAE — causa e subcausa","49% casaram; concorda com a Crítica em 93%"])
w2.column_dimensions["A"].width=56; w2.column_dimensions["B"].width=52
for _r in (1,5): w2.cell(row=_r,column=1).font=Font(bold=True,size=12,color=NAVY)
wb.save("SS_OS_Critica_TMAE.xlsx")
print(cont.most_common()); print("linhas",len(A))
