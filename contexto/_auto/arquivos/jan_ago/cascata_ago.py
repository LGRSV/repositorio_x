# -*- coding: utf-8 -*-
import json, collections, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
d=json.load(open("as_1696.json")); J={x["ss"]:x for x in json.load(open("cr_tm_por_ss.json"))}
TM={x["ss"]:x for x in json.load(open("agosto_tmae.json"))}
base1671={x["ss"] for x in json.load(open("confronto_origem.json"))}
def U(r,c): return str(r.get(c) or "").strip().upper()
def limpo(v):
    s=str(v or "").strip().upper(); return "" if s in ("","0","00","000","00000","N/A","NA","-","NONE") else s
AGO=[r for r in d if str(r.get("mes"))=="2026-08" and r["ss"] in base1671]
print("agosto dentro das 1.671:",len(AGO))
GAT={"FURTO":"furto","REMANEJAMENTO":"remanejamento","SEM TROCA (NÃO SUBSTITUÍDO)":"sem_troca",
 "DIVISÃO DE CIRCUITO":"divisao","MELHORIA DE POSTO":"melhoria","INCONCLUSIVO":"inconclusivo",
 "SEM OS E SEM OBRA":"sem_os","PREVENTIVO/PROGRAMADO":"preventivo","ABALROAMENTO":"abalroamento",
 "TAPE/REGULARIZAÇÃO DE TENSÃO":"tap","AUXILIAR DE RELIGADOR":"auxiliar","AUSENTE DA CRÍTICA":"sem_interrupcao",
 "FORA DA JANELA DA CRÍTICA":"fora_da_janela","ERRO DE CADASTRO — NÃO É TRANSFORMADOR":"erro_cadastro",
 "DANO DE TERCEIROS":"terceiros","FALTA DE FASE":"falta_fase","SS DUPLICADA":"duplicada",
 "OBRA SEM TRANSFORMADOR NO MATERIAL":"obra_sem_transformador","OBRA SEM EXECUÇÃO":"obra_sem_execucao",
 "AVALIAR COM O MATHEUS":"avaliar_matheus","SEM OBRA — PASSADOS 60 DIAS":"sem_obra"}
linhas=[]
for r in AGO:
    j=J.get(r["ss"],{}); t=TM.get(r["ss"],{})
    cat=U(r,"categoria"); conta=U(r,"conta")=="SIM"
    temcr=bool(j.get("cr_causa"))
    troca = limpo(r.get("ns_ret")) and limpo(r.get("ns_inst")) and limpo(r.get("ns_ret"))!=limpo(r.get("ns_inst"))
    if not conta:
        res="EXCLUÍDA"; gat=GAT.get(cat,"outro")
    elif not temcr:
        if troca: res="RETIDO — SEM INTERRUPÇÃO NA JANELA"; gat="sem_interrupcao"
        else: res="EXCLUÍDA"; gat="sem_interrupcao"
    else:
        res="SAÍDA"; gat=""
    linhas.append({"ss":r["ss"],"trafo":r.get("trafo"),"abertura":str(r.get("abertura"))[:16],
      "cat":cat,"conta":"SIM" if conta else "NÃO","resultado":res,"gatilho":gat,
      "critica":"sim" if temcr else "NÃO","cr_causa":j.get("cr_causa",""),"cr_sub":j.get("cr_sub",""),
      "tmae_janela":t.get("na_janela",0),"tmae_ativo":t.get("no_ativo",0),
      "tmae_causa":t.get("causa",""),"tmae_sub":t.get("sub",""),"tmae_obs":t.get("obs",""),
      "ns_ret":r.get("ns_ret"),"ns_inst":r.get("ns_inst"),"troca":"sim" if troca else "não",
      "equipe":r.get("equipe"),"tipo":r.get("tipo_ss"),"motivo":r.get("motivo","")})
c=collections.Counter(x["resultado"] for x in linhas)
print(c.most_common())
print("gatilhos das EXCLUÍDA:",collections.Counter(x["gatilho"] for x in linhas if x["resultado"]=="EXCLUÍDA").most_common())
NAVY="1F3864"; borda=Border(*[Side(style="thin",color="D9D9D9")]*4)
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Agosto — cascata"
cols=[("SS","ss"),("Trafo","trafo"),("Abertura","abertura"),("Tipo da SS","tipo"),("Equipe","equipe"),
 ("Categoria","cat"),("Conta","conta"),("Resultado","resultado"),("Gatilho","gatilho"),
 ("Casou na Crítica","critica"),("Crítica — causa","cr_causa"),("Crítica — subcausa","cr_sub"),
 ("TMAE na janela","tmae_janela"),("TMAE no ativo","tmae_ativo"),("TMAE — causa","tmae_causa"),
 ("TMAE — subcausa","tmae_sub"),("TMAE — observação","tmae_obs"),
 ("NS retirado","ns_ret"),("NS instalado","ns_inst"),("Prova de troca","troca"),("Motivo da leitura","motivo")]
ws.append([c for c,_ in cols])
for c_ in ws[1]:
    c_.font=Font(bold=True,color="FFFFFF",size=10); c_.fill=PatternFill("solid",fgColor=NAVY)
    c_.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
ordem={"EXCLUÍDA":0,"RETIDO — SEM INTERRUPÇÃO NA JANELA":1,"SAÍDA":2}
for x in sorted(linhas,key=lambda x:(ordem[x["resultado"]],x["ss"])):
    ws.append([x[k] for _,k in cols])
    cor={"EXCLUÍDA":"FCE4E4","RETIDO — SEM INTERRUPÇÃO NA JANELA":"FFF2CC","SAÍDA":"E2EFDA"}[x["resultado"]]
    ws.cell(row=ws.max_row,column=8).fill=PatternFill("solid",fgColor=cor)
for i,(c_,_) in enumerate(cols,1):
    ws.column_dimensions[get_column_letter(i)].width={"Motivo da leitura":50,"TMAE — observação":45,
      "Crítica — subcausa":32,"TMAE — subcausa":30,"Resultado":32,"Categoria":26,"SS":22,"Tipo da SS":26}.get(c_,13)
for row in ws.iter_rows(min_row=2):
    for c_ in row: c_.border=borda; c_.alignment=Alignment(vertical="top",wrap_text=c_.column in (17,21))
ws.freeze_panes="B2"; ws.auto_filter.ref=ws.dimensions; ws.row_dimensions[1].height=36
w2=wb.create_sheet("Fecho",0)
for t in [["Expurgos de agosto — com a Crítica e o TMAE de agosto aplicados",""],["",""],
 ["SS de agosto nas 1.671",len(AGO)],
 ["EXCLUÍDA (expurgo)",c["EXCLUÍDA"]],
 ["RETIDO — sem interrupção mas com prova de troca",c["RETIDO — SEM INTERRUPÇÃO NA JANELA"]],
 ["SAÍDA (contam)",c["SAÍDA"]],["",""],
 ["No site, até julho",220],
 ["Menos os 16 de sem obra (temporais)",204],
 ["Com agosto: 204 + %d"%c["EXCLUÍDA"],204+c["EXCLUÍDA"]],
 ["Com agosto, contando os 16: 220 + %d"%c["EXCLUÍDA"],220+c["EXCLUÍDA"]]]:
    w2.append(t)
w2["A1"].font=Font(bold=True,size=13,color=NAVY)
w2.column_dimensions["A"].width=52; w2.column_dimensions["B"].width=12
wb.save("Agosto_cascata_expurgos.xlsx"); print("xlsx ok")
