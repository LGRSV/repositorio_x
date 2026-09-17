# -*- coding: utf-8 -*-
import json, csv, collections, re
from openpyxl import load_workbook
csv.field_size_limit(10**8)
B="/home/user/repositorio_x/auditoria-transformadores-134/public/bases/originais/"

def nobra(v):
    s=re.sub(r"\D","",str(v or ""))
    return s.lstrip("0") or ""

# 1) base SS/OS 64 colunas
wb=load_workbook(B+"Original_OS.xlsx",read_only=True)
ws=wb["BASE SS_OS"]; it=ws.iter_rows(values_only=True); h=[str(c).strip() for c in next(it)]
ssos={}; portrafo=collections.defaultdict(list)
for row in it:
    r={h[i]:row[i] for i in range(min(len(h),len(row)))}
    k=str(r.get("NUMERO_SS") or "").strip()
    if k: ssos[k]=r
    t=str(r.get("NUM_TRAFO") or "").strip()
    if t: portrafo[t].append(r)
wb.close()
print("SS/OS",len(ssos),"trafos",len(portrafo))

# 2) obras
wb=load_workbook(B+"Original_Cadastro_Obras.xlsx",read_only=True)
ws=wb["Export"]; it=ws.iter_rows(values_only=True); ho=[str(c).strip() for c in next(it)]
alvo=["NUM_OBRA","NUM_OS","DSC_STATUS","DSC_OCORRENCIA","DESCRICAO_OBRA","TIPO_OBRA","DTH_ABERTURA","DTH_ENCERRAMENTO","DTH_INICIO_FISICO","DTH_TERMINO_FISICO","DATA_CONCLUSAO_FISICA","MT_REALIZADO","TOTAL_REALIZADO","MOT_SUSPENSAO","REGIONAL"]
ix={c:ho.index(c) for c in alvo if c in ho}
obras={}
for row in it:
    k=nobra(row[ix["NUM_OBRA"]])
    if k: obras[k]={c:(str(row[i]) if row[i] is not None else "") for c,i in ix.items()}
wb.close()
print("obras",len(obras))

# 3) material
mat=collections.defaultdict(list)
for f in ["df9ff469-material_obra.txt","71743a90-material_obra_1295.txt"]:
    for r in csv.DictReader(open(f,encoding="latin-1"),delimiter="\t"):
        k=nobra(r.get("num_obra"))
        if k: mat[k].append(r)
print("obras com material",len(mat))

# 4) TMAE por ativo
tmae=collections.defaultdict(list)
with open("32196f1a-TMAE_2026_Jan_Jun_Consolidado.txt",encoding="latin-1") as f:
    for r in csv.DictReader(f,delimiter=";"):
        for c in ("COD_ELE_REDE_TNT","COD_INS_TRF_TNT"):
            v=str(r.get(c) or "").strip()
            if v and v not in ("0",""): tmae[v].append(r)
print("ativos no TMAE",len(tmae))

json.dump({"obras":obras},open("obras.json","w"))
json.dump({k:[{c:v[c] for c in ("cod_material","descricao","qtdprevista","qtdrealizada","val_requisitado")} for v in vs] for k,vs in mat.items()},open("material.json","w"))
# tmae: guarda só campos úteis
red={}
for k,vs in tmae.items():
    red[k]=[{c:str(v.get(c) or "") for c in ("DTA_ORIG_TNT","DTA_CONCL_TNT","DES_CAUSA_INTER_CAU","DES_SUB_CAUSA_INTER_SCR","GRUPO_CAUSA","EQUIPE","DES_OBS_EXEC_SERV_TNT","COD_TIPO_CONCL_TNT","NOM_LOC_TNT","DES_ALIM_TNT")} for v in vs]
json.dump(red,open("tmae.json","w"))
json.dump({k:{c:(str(v[c]) if v.get(c) is not None else "") for c in h} for k,v in ssos.items()},open("ssos.json","w"))
json.dump({k:[str(x.get("NUMERO_SS") or "") for x in v] for k,v in portrafo.items()},open("portrafo.json","w"))
print("ok")
