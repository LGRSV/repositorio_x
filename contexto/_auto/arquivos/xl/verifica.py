import openpyxl, re, warnings; warnings.filterwarnings('ignore')
from openpyxl.utils import column_index_from_string as ci
wb=openpyxl.load_workbook('Expurgos_Jan_Ago_2026.xlsx')
data=[r for r in wb['Expurgos'].iter_rows(min_row=2,max_row=228,values_only=True)]
def col(letter): return [r[ci(letter)-1] for r in data]

def countifs(args):
    pares=[(col(re.match(r"Expurgos!\$([A-Z]+)\$\d+:\$[A-Z]+\$\d+",a).group(1)), c.strip('"'))
           for a,c in zip(args[0::2],args[1::2])]
    n=0
    for i in range(len(data)):
        if all(str(v[i] if v[i] is not None else '')==c for v,c in pares): n+=1
    return n
def counta(args):
    c=col(re.match(r"Expurgos!\$([A-Z]+)\$\d+:\$[A-Z]+\$\d+",args[0]).group(1))
    return sum(1 for v in c if v not in (None,''))

def split_args(s):
    out=[];d=0;cur=''
    for ch in s:
        if ch=='(':d+=1
        if ch==')':d-=1
        if ch==',' and d==0: out.append(cur.strip()); cur=''
        else: cur+=ch
    if cur.strip(): out.append(cur.strip())
    return out

ws=wb['Resumo']
cells={}
for row in ws.iter_rows(min_row=1,max_row=45,max_col=5):
    for c in row:
        if isinstance(c.value,str) and c.value.startswith('='):
            cells[c.coordinate]=c.value

def ev(coord,seen=None):
    seen=seen or set()
    if coord in seen: raise RuntimeError('ciclo em '+coord)
    f=cells.get(coord)
    if f is None:
        v=ws[coord].value
        return v if isinstance(v,(int,float)) else 0
    e=f[1:]
    for fn,h in (('COUNTIFS',countifs),('COUNTA',counta)):
        while True:
            m=re.search(fn+r'\(',e)
            if not m: break
            i=m.end(); d=1; j=i
            while d>0: 
                if e[j]=='(': d+=1
                elif e[j]==')': d-=1
                j+=1
            e=e[:m.start()]+str(h(split_args(e[i:j-1])))+e[j:]
    while True:
        m=re.search(r'SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)',e)
        if not m: break
        tot=sum(ev(f'{m.group(1)}{r}',seen|{coord}) for r in range(int(m.group(2)),int(m.group(4))+1))
        e=e[:m.start()]+str(tot)+e[m.end():]
    m=re.match(r'^IF\((\w+)=0,0,(\w+)/(\w+)\)$',e)
    if m:
        den=ev(m.group(1),seen|{coord})
        return 0 if den==0 else ev(m.group(2),seen|{coord})/ev(m.group(3),seen|{coord})
    e=re.sub(r'\$','',e)
    for ref in sorted(set(re.findall(r'\b([A-Z]{1,2}\d{1,3})\b',e)),key=len,reverse=True):
        e=re.sub(r'\b'+ref+r'\b',str(ev(ref,seen|{coord})),e)
    return eval(e)

ESP={'B5':115,'B6':112,'B7':227,'B11':112,'B12':49,'B13':38,'B14':16,'B15':12,'B16':227,
     'B20':112,'B21':60,'B22':172,'B39':81,'B40':31,'B41':112,
     'B26':47,'C26':23,'B27':35,'C27':24,'B28':34,'C28':24,'B29':31,'C29':19,
     'B30':16,'C30':7,'B31':28,'C31':9,'B32':13,'C32':0,'B33':23,'C33':6,'B34':227,'C34':112}
ok=bad=0
for k,esp in sorted(ESP.items(), key=lambda x:(x[0][0],int(x[0][1:]))):
    try: got=ev(k)
    except Exception as ex: print(f'  {k}: ERRO {ex}'); bad+=1; continue
    st='OK ' if got==esp else 'ERRADO'
    if got==esp: ok+=1
    else: bad+=1; print(f'  {k}: {st} esperado {esp}, formula deu {got}   <<< {cells.get(k)}')
print(f'\n{ok} formulas conferem, {bad} divergem, de {len(ESP)} verificadas')
# percentuais
for k in ['C5','C6','C11','C20','C39','D26','D32']:
    if k in cells:
        try: print(f'  {k} = {ev(k):.3f}   ({cells[k]})')
        except Exception as ex: print(f'  {k} ERRO {ex}')
