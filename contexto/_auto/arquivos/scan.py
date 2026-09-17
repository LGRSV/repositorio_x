import openpyxl, glob, os, warnings
warnings.filterwarnings('ignore')
files = sorted(glob.glob('*.xlsx')) + sorted(glob.glob('originais/*.xlsx'))
for f in files:
    try:
        wb = openpyxl.load_workbook(f, read_only=True)
    except Exception as e:
        print(f'{f}: ERRO {e}'); continue
    for sh in wb.sheetnames:
        ws = wb[sh]
        it = ws.iter_rows(values_only=True)
        try: hdr = list(next(it))
        except StopIteration: continue
        hdr = [str(h) if h is not None else '' for h in hdr]
        obracols = [h for h in hdr if 'OBRA' in h.upper()]
        if not obracols: continue
        # conta obras distintas na coluna principal
        key = 'NUM_OBRA' if 'NUM_OBRA' in hdr else obracols[0]
        ki = hdr.index(key)
        vals, n = set(), 0
        for r in it:
            n += 1
            v = r[ki] if ki < len(r) else None
            if v not in (None, ''): vals.add(str(v).strip())
        print(f'{f} [{sh}]  linhas={n}  col={key}  obras_distintas={len(vals)}  outras_cols_obra={obracols[:6]}')
    wb.close()
