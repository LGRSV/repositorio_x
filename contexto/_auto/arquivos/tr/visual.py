import json, collections, datetime, re, warnings; warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gl
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule, CellIsRule
d=json.load(open('dossie.json'))
MES={'01':'jan','02':'fev','03':'mar','04':'abr','05':'mai','06':'jun','07':'jul','08':'ago'}
for o in d: o['mes']=MES.get((o['abertura'] or '')[5:7],'')
d=sorted(d,key=lambda o:o['abertura'] or '9')
F='Calibri'
NAVY='1F1F4C'; BURN='C55A11'; SURGE='2F5597'; OK='548235'; LATE='C00000'
ZEB='F4F6F9'; LINE='D9D9D9'; MUT='595959'
fill=lambda c: PatternFill('solid',fgColor=c)
B=Side(style='thin',color=LINE); BOX=Border(left=B,right=B,top=B,bottom=B)
H=Font(name=F,size=11,bold=True,color='FFFFFF')
def cab(ws,row,cols,widths=None,cor=NAVY):
    for j,c in enumerate(cols,1):
        x=ws.cell(row,j,c); x.font=H; x.fill=fill(cor); x.border=BOX
        x.alignment=Alignment(vertical='center',wrap_text=True,horizontal='center' if j>1 else 'left')
    ws.row_dimensions[row].height=30
    if widths:
        for j,w in enumerate(widths,1): ws.column_dimensions[gl(j)].width=w
def tit(ws,t,s=None):
    ws['A1']=t; ws['A1'].font=Font(name=F,size=17,bold=True,color=NAVY); ws.row_dimensions[1].height=24
    if s: ws['A2']=s; ws['A2'].font=Font(name=F,size=10.5,color=MUT)
    ws.sheet_view.showGridLines=False

wb=openpyxl.Workbook()
# ================= PAINEL =================
ws=wb.active; ws.title='Painel'
tit(ws,'Por que queimaram — 16 trafos rurais de 112,5 e 150 kVA',
   'Janeiro a agosto de 2026. Causa registrada na Crítica, laudo de campo e cumprimento do prazo de 72 horas.')
sub=collections.Counter(o['oc_sub'] for o in d if o['oc_sub'])
camp=collections.Counter(o['motivo_campo'] for o in d if o['motivo_campo'])
meses=collections.Counter(o['mes'] for o in d if o['mes'])
noprazo=[o for o in d if o['h_prazo'] is not None and o['h_prazo']>=0]
atraso=[o for o in d if o['h_prazo'] is not None and o['h_prazo']<0]
# indicadores
ws['A4']='O QUE O PERÍODO MOSTRA'; ws['A4'].font=Font(name=F,size=11,bold=True,color=MUT)
ind=[('Transformadores analisados',len(d),NAVY,''),
     ('Com interrupção registrada na Crítica',sum(1 for o in d if o['oc_sub']),SURGE,'de 16'),
     ('Queimados por descarga atmosférica',sum(1 for o in d if 'DESCARGA' in (o['oc_sub'] or '')),BURN,'subcausa dominante'),
     ('Laudos de campo sem causa definida',camp.get('NÃO IDENTIFICADO',0),BURN,'quase metade'),
     ('Atendimentos dentro das 72 h',len(noprazo),OK,f'de {len(noprazo)+len(atraso)}'),
     ('Fora do prazo',len(atraso),LATE,'ETO-RD-DP 00011/2026')]
r=5
for lab,val,cor,nota in ind:
    ws.cell(r,1,lab).font=Font(name=F,size=11)
    c=ws.cell(r,2,val); c.font=Font(name=F,size=20,bold=True,color=cor); c.alignment=Alignment(horizontal='center')
    ws.cell(r,3,nota).font=Font(name=F,size=10,italic=True,color=MUT)
    ws.row_dimensions[r].height=26; r+=1
for j,w in enumerate([44,10,34,3,30,10,3,30,10],1): ws.column_dimensions[gl(j)].width=w
# tabelas de apoio para os graficos
ws['E4']='SUBCAUSA NA CRÍTICA'; ws['E4'].font=Font(name=F,size=11,bold=True,color=SURGE)
r=5
for k,v in sub.most_common():
    ws.cell(r,5,k.title()).font=Font(name=F,size=10.5); ws.cell(r,6,v).font=Font(name=F,size=10.5); r+=1
fim_sub=r-1
ws['H4']='PROVÁVEL MOTIVO EM CAMPO'; ws['H4'].font=Font(name=F,size=11,bold=True,color=BURN)
r=5
for k,v in camp.most_common():
    ws.cell(r,8,k.title()).font=Font(name=F,size=10.5); ws.cell(r,9,v).font=Font(name=F,size=10.5); r+=1
fim_camp=r-1
def grafico(ws,tit_,ref_cat,ref_val,cor,anchor,h=7.2):
    ch=BarChart(); ch.type='bar'; ch.style=2; ch.title=tit_; ch.height=h; ch.width=13
    ch.add_data(Reference(ws,min_col=ref_val[0],min_row=ref_val[1],max_row=ref_val[2]),titles_from_data=False)
    ch.set_categories(Reference(ws,min_col=ref_cat[0],min_row=ref_cat[1],max_row=ref_cat[2]))
    ch.legend=None; ch.dataLabels=None
    from openpyxl.chart.label import DataLabelList
    ch.dLbls=DataLabelList(); ch.dLbls.showVal=True
    ch.series[0].graphicalProperties.solidFill=cor
    ch.series[0].graphicalProperties.line.noFill=True
    ch.y_axis.majorGridlines=None; ch.x_axis.delete=False; ch.y_axis.delete=False
    ws.add_chart(ch,anchor)
grafico(ws,'Subcausa registrada na Crítica',(5,5,fim_sub),(6,5,fim_sub),SURGE,'A13')
grafico(ws,'Provável motivo apontado em campo',(8,5,fim_camp),(9,5,fim_camp),BURN,'E13')
mr=5
ws['K4']='MÊS'; 
for k in ['jan','fev','mar','abr','mai','jun','jul','ago']:
    ws.cell(mr,11,k); ws.cell(mr,12,meses.get(k,0)); mr+=1
grafico(ws,'Quando abriu a SS',(11,5,mr-1),(12,5,mr-1),NAVY,'A29',h=6.5)
ws.column_dimensions['K'].width=8; ws.column_dimensions['L'].width=6
ws['A45']='Duas datas não fecham: em ETO-RD-GU 00061/2026 a ocorrência casada é de 30/04 e a SS é de 13/01; em ETO-RD-GR 00405/2026 a SS abriu 5 h antes da interrupção. Conferir antes de usar a causa.'
ws['A45'].font=Font(name=F,size=10.5,italic=True,color=BURN)

# ================= LINHA DO TEMPO =================
ws2=wb.create_sheet('Linha do tempo')
tit(ws2,'Cada caso contra o prazo de 72 horas',
   'Escala comum de −60 h a +110 h em torno da abertura da SS. Azul = interrupção na Crítica · laranja = SS aberta · cinza = prazo de 72 h · verde/vermelho = conclusão.')
H0,H1,STEP=-60,110,5
NC=(H1-H0)//STEP
COL0=6
cab(ws2,4,['Transformador','SS','Local','Atend. (h)','Folga (h)']+['']*NC,[16,21,20,11,11]+[2.1]*NC)
for k in range(NC):
    h=H0+k*STEP
    c=ws2.cell(4,COL0+k,h if h%30==0 else None)
    c.font=Font(name=F,size=8,bold=True,color='FFFFFF'); c.alignment=Alignment(horizontal='center')
r=5
for o in d:
    ws2.cell(r,1,o['trafo']).font=Font(name=F,size=10.5,bold=True)
    ws2.cell(r,2,o['ss']).font=Font(name=F,size=10)
    ws2.cell(r,3,o['loc']).font=Font(name=F,size=10)
    ws2.cell(r,4,round(o['h_atend'],1) if o['h_atend'] is not None else None).font=Font(name=F,size=10)
    fc=ws2.cell(r,5,round(o['h_prazo'],1) if o['h_prazo'] is not None else None)
    fc.font=Font(name=F,size=10,bold=True,color=(LATE if (o['h_prazo'] or 0)<0 else OK))
    for c in range(1,6): ws2.cell(r,c).border=BOX; ws2.cell(r,c).alignment=Alignment(vertical='center',horizontal='center' if c>3 else 'left')
    idx=lambda h: int(round((h-H0)/STEP))
    for k in range(NC): ws2.cell(r,COL0+k).fill=fill(ZEB if r%2 else 'FFFFFF')
    if o['oc_ini'] and o['h_oc_ss'] is not None and o['h_oc_ss']>-100:
        a=idx(-o['h_oc_ss']); b=idx(-o['h_oc_ss']+float(o['oc_dur'] or 0))
        for k in range(max(0,a),min(NC,max(b,a+1))): ws2.cell(r,COL0+k).fill=fill(SURGE)
    if 0<=idx(0)<NC: ws2.cell(r,COL0+idx(0)).fill=fill(BURN)
    if 0<=idx(72)<NC: ws2.cell(r,COL0+idx(72)).fill=fill('BFBFBF')
    if o['h_atend'] is not None and 0<=idx(o['h_atend'])<NC:
        ws2.cell(r,COL0+idx(o['h_atend'])).fill=fill(LATE if (o['h_prazo'] or 0)<0 else OK)
    ws2.row_dimensions[r].height=17; r+=1
ws2.conditional_formatting.add(f'D5:D{r-1}',DataBarRule(start_type='num',start_value=0,end_type='num',end_value=110,color=BURN))
ws2.conditional_formatting.add(f'E5:E{r-1}',CellIsRule(operator='lessThan',formula=['0'],fill=fill('FCE4E4')))
ws2.freeze_panes='F5'
ws2.cell(r+1,1,'Legenda:').font=Font(name=F,size=10,bold=True,color=MUT)
for j,(lab,cor) in enumerate([('interrupção',SURGE),('SS aberta',BURN),('prazo 72 h','BFBFBF'),('concluída no prazo',OK),('concluída em atraso',LATE)]):
    ws2.cell(r+1,2+j*2).fill=fill(cor); ws2.cell(r+1,3+j*2,lab).font=Font(name=F,size=9.5,color=MUT)

# ================= CAUSA =================
ws3=wb.create_sheet('Causa')
tit(ws3,'O que a Crítica registrou × o que o campo encontrou',
   'Linha destacada quando as duas fontes não apontam a mesma coisa.')
cab(ws3,4,['Transformador','SS','Mês','Crítica — causa','Crítica — subcausa','Campo — provável motivo','Batem?','Para-raios','Elo','Aterramento Ω','Potência'],
    [16,21,7,24,30,24,10,11,7,18,16])
r=5
for o in d:
    A='DESCARGA' in (o['oc_sub'] or ''); Bc='DESCARGA' in (o['motivo_campo'] or '')
    bate='—'
    if o['oc_sub'] and o['motivo_campo']:
        bate='sim' if (A==Bc and o['motivo_campo']!='NÃO IDENTIFICADO') else 'não'
    at=' · '.join(str(x).replace('.',',') for x in o['at'] if x not in ('','None','0')) or '—'
    pt=(f"{o['pot_ret']:g} → {o['pot_inst']:g} kVA" if o['pot_ret'] and o['pot_inst'] else '—')
    vals=[o['trafo'],o['ss'],o['mes'],o['oc_causa'] or '—',o['oc_sub'] or 'sem interrupção registrada',
          o['motivo_campo'] or 'não preenchido',bate,o['para_raios'] or '—',o['elo'] or '—',at,pt]
    for j,v in enumerate(vals,1):
        c=ws3.cell(r,j,v); c.border=BOX; c.font=Font(name=F,size=10.5)
        c.alignment=Alignment(vertical='center',wrap_text=(j in(4,5,6)),horizontal='center' if j in(3,7,8,9) else 'left')
        if r%2==0: c.fill=fill(ZEB)
    if bate=='não':
        for j in range(1,12): ws3.cell(r,j).fill=fill('FDECE4')
        ws3.cell(r,7).font=Font(name=F,size=10.5,bold=True,color=BURN)
    elif bate=='sim': ws3.cell(r,7).font=Font(name=F,size=10.5,bold=True,color=OK)
    if o['pot_ret'] and o['pot_inst'] and o['pot_ret']!=o['pot_inst']:
        ws3.cell(r,11).font=Font(name=F,size=10.5,bold=True,color=(BURN if o['pot_inst']>o['pot_ret'] else SURGE))
    ws3.row_dimensions[r].height=30; r+=1
ws3.freeze_panes='A5'; ws3.auto_filter.ref=f'A4:K{r-1}'

# ================= DOSSIE =================
ws4=wb.create_sheet('Dossiê')
tit(ws4,'Ficha por transformador','Cronologia completa, equipamento retirado e instalado, e os textos da SS e da OS.')
cab(ws4,4,['Transformador','SS','Local','Clientes','Interrupção início','Interrupção fim','Dur. (h)','SS aberta','Prazo limite','SS terminada','Atend. (h)','Folga (h)','Série retirada','Série instalada','Fabricante ret.','Fabricante inst.','Texto da SS','Texto da OS'],
    [16,21,18,9,16,16,9,16,16,16,10,10,14,14,15,15,60,60])
r=5
dtbr=lambda s: (datetime.datetime.strptime(s[:16],'%Y-%m-%d %H:%M').strftime('%d/%m/%Y %H:%M') if s else '—')
for o in d:
    vals=[o['trafo'],o['ss'],o['loc'],o['clientes'] or '0',dtbr(o['oc_ini']),dtbr(o['oc_fim']),
          o['oc_dur'] or '—',dtbr(o['abertura']),dtbr(o['limite']),dtbr(o['termino']),
          round(o['h_atend'],1) if o['h_atend'] is not None else None,
          round(o['h_prazo'],1) if o['h_prazo'] is not None else None,
          o['ns_ret'] or '—',o['ns_inst'] or '—',o['fab_ret'] or '—',o['fab_inst'] or '—',
          re.sub(r'\s+',' ',(o['texto_ss'] or '').replace('_x000D_',' ')),
          re.sub(r'\s+',' ',(o['texto_os'] or '').replace('_x000D_',' '))]
    for j,v in enumerate(vals,1):
        c=ws4.cell(r,j,v); c.border=BOX; c.font=Font(name=F,size=10)
        c.alignment=Alignment(vertical='top',wrap_text=(j>=17))
        if r%2==0: c.fill=fill(ZEB)
    ws4.cell(r,12).font=Font(name=F,size=10,bold=True,color=(LATE if (o['h_prazo'] or 0)<0 else OK))
    ws4.row_dimensions[r].height=46; r+=1
ws4.freeze_panes='C5'; ws4.auto_filter.ref=f'A4:R{r-1}'
wb.save('Trafos_rurais_visual.xlsx')
print('salvo | abas:',wb.sheetnames)
