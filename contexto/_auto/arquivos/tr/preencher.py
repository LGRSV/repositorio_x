import openpyxl, json, gzip, re, datetime, collections, warnings; warnings.filterwarnings('ignore')
U='/root/.claude/uploads/74dc9c64-5026-54ee-a81e-173d2f38a735/'
B='/home/user/repositorio_x/auditoria-transformadores-134/'
def g(r,i): return r[i] if (i is not None and i<len(r)) else None
def cheia(r): return sum(1 for v in r if v not in (None,''))
def txt(v):
    if v in (None,''): return ''
    if isinstance(v,datetime.datetime): return v.strftime('%d/%m/%Y')
    s=str(v).strip()
    s=re.sub(r'\s+\d{2}:\d{2}:\d{2}$','',s)
    if s.endswith('.0'): s=s[:-2]
    return s
def pot(v):
    s=txt(v)
    if not s: return ''
    try: n=float(s.replace(',','.'))
    except: return s
    if n==1125: n=112.5
    return (f'{n:g}'.replace('.',','))+' kVA'
def num(v):
    s=txt(v)
    return s.replace('.',',') if re.match(r'^\d+\.\d+$',s) else s
def med(r,IS,n):
    dep=[g(r,IS.get(f'MEDIÇÃO_DE_ATERRAMENTO_X{k}')) for k in (1,2,3)]
    ant=[g(r,IS.get(f'MEDIÇÃO_DE_ATERRAMENTO_ANTES_X{k}')) for k in (1,2,3)]
    def ok(L):
        try: return any(float(str(x).replace(',','.'))!=0 for x in L if x not in (None,''))
        except: return any(x not in (None,'') for x in L)
    src=dep if ok(dep) else (ant if ok(ant) else None)
    if src is None: return ''
    return num(src[n-1])
def ref(v):
    s=txt(v)
    return '-' if s.upper() in ('NÃO É REFORMADO','NAO É REFORMADO','NÃO E REFORMADO') else s

# fontes
w=openpyxl.load_workbook(U+'8cfc359b-Trafo.xlsx',read_only=True); s=w['BASE_SS_OS']
i=s.iter_rows(values_only=True); h=[' '.join(str(x).split()) if x is not None else '' for x in next(i)]
IS={k:j for j,k in enumerate(h)}; best={}
for r in i:
    k=str(g(r,IS['NUMERO_SS']) or '').strip()
    if k and (k not in best or cheia(r)>cheia(best[k])): best[k]=r
w.close()
site={x['ss']:x for x in json.load(open(B+'public/fluxo-1582.json'))['registros']}
fis=json.load(gzip.open(B+'dados/fis-2026-07-energisa.json.gz'))

TPL={
 23:['Realizado reparo cola e fita','Poste retirado','CITD - Retirado','N° de série','Potência','TAP','Tensão','Fabricante','Data de Fabricação','','Reformado','Reformadora','Data de Reforma'],
 24:['Realizado reparo cola e fita','Poste instalado','CITD - instalado','N° de série','Potência','TAP','Tensão','Fabricante','Data de Fabricação','','Reformado','Reformadora','Data de Reforma'],
 25:['Feito melhoria no aterramento','Medição de aterramento X1','Medição de aterramento X2','Medição de aterramento X3','Cabo de aterramento está integro?','Barramento está em boas condições'],
 26:['Para raios','Equipamento apresenta vazamento de óleo','Furto de energia no local?','Vegetação tocando a rede?','Provável motivo','Elo instalado','Número da intervenção'],
}
def monta(idx,vals,fim_nl=True):
    ls=[]
    for lab in TPL[idx]:
        if lab=='': ls.append(''); continue
        ls.append(f'{lab}: {vals.get(lab,"")}'.rstrip() if vals.get(lab,'')=='' else f'{lab}: {vals[lab]}')
    t='\n'.join(ls)
    return t+'\n' if fim_nl else t
def bare(v,idx):
    if v in (None,''): return True
    t=str(v)
    for lab in TPL[idx]:
        if lab: t=t.replace(lab+':','')
    return not re.sub(r'[\s\n]','',t)

wb=openpyxl.load_workbook(U+'ec33c0b4-analise_trafos_rurais_1125_e_150.xlsx')
ws=wb['Trafos 112,5 e 150 rural']
preenchidos=collections.Counter()
for rr in range(2,ws.max_row+1):
    ss=str(ws.cell(rr,12).value or '').strip()
    if not ss: continue
    r=best.get(ss); sv=site.get(ss)
    cod=re.findall(r'\d{10}',str(ws.cell(rr,5).value or ''))
    f=fis.get(cod[0]) if cod else None
    tensao=(f'{f["TENSAO_PRIMARIA"]}'.replace('.',',')+' kV') if f and f.get('TENSAO_PRIMARIA') else ''
    V=lambda c: txt(g(r,IS.get(c))) if r else ''
    def put(col,val):
        if val in (None,''): return
        atual=ws.cell(rr,col).value
        if atual in (None,'') or (col in TPL and bare(atual,col)):
            ws.cell(rr,col).value=val; preenchidos[col]+=1
    if r:
        put(18,ws.cell(rr,18).value or V('NUMERO_OS'))
        put(19,V('ORIGEM')); put(20,V('DEFEITO'))
        put(22,' '.join(str(g(r,IS['DESCRICAO_OS']) or '').split()))
        put(27,V('NUM_OBRA'))
        put(23,monta(23,{'Realizado reparo cola e fita':V('FEITO_COLA_E_FITA'),'CITD - Retirado':V('TOMBAMENTO_RETIRADO'),
           'N° de série':V('NS_RETIRADO'),'Potência':pot(g(r,IS['POTENCIA_RET'])),'Tensão':tensao,
           'Fabricante':V('FABRICANTE_RETIRADO'),'Data de Fabricação':V('DATA_FABRICACAO_RETIRADO'),
           'Reformado':V('TRANSFORMADOR_RETIRADO_É_REFORMADO'),'Reformadora':ref(g(r,IS.get('REFORMADORA_DESAT'))),
           'Data de Reforma':V('DATA_REFORMA_RET')}))
        put(24,monta(24,{'Realizado reparo cola e fita':V('FEITO_COLA_E_FITA'),'CITD - instalado':V('TOMBAMENTO_INSTALADO'),
           'N° de série':V('NS_INSTALADO'),'Potência':pot(g(r,IS['POTENCIA_INST'])),'Tensão':tensao,
           'Fabricante':V('FABRICANTE_INSTALADO'),'Data de Fabricação':V('DATA_FABRICACAO_INSTALADO'),
           'Reformado':V('TRAFO_INSTALADO_É_REFORMADO'),'Reformadora':ref(g(r,IS.get('REFORMADORA_INSTALADO'))),
           'Data de Reforma':V('DATA_REFORMA_INST')}))
        put(25,monta(25,{'Feito melhoria no aterramento':V('FEZ_MELHORIA_DE_ATERRAMENTO'),
           'Medição de aterramento X1':med(r,IS,1),
           'Medição de aterramento X2':med(r,IS,2),
           'Medição de aterramento X3':med(r,IS,3)},fim_nl=False)
           +(f"\nAterramento ligado ao tanque: {V('ATERRAMENTO_ESTA_CONECTADO_AO_TANQUE_DO_TRAFO')}" if V('ATERRAMENTO_ESTA_CONECTADO_AO_TANQUE_DO_TRAFO') else '')
           +(f"\nTipo de barramento: {V('TIPO_DE_BARRAMENTO')}" if V('TIPO_DE_BARRAMENTO') else ''))
        put(26,monta(26,{'Para raios':V('PARA_RAIOS_INSTALADO_SUBSTITUIDO'),
           'Equipamento apresenta vazamento de óleo':V('EQUIPAMENTO_APRESENTA_VAZAMENTO_DE_OLEO'),
           'Vegetação tocando a rede?':V('EXISTE_VEGETAÇÃO_TOCANDO_A_REDE_BT_OU_MT'),
           'Provável motivo':V('PROVÁVEL_MOTIVO_DO_DEFEITO'),'Elo instalado':V('QUAL_ELO_INSTALADO_NO_LOCAL'),
           'Número da intervenção':V('NUM_INT')},fim_nl=False))
    # Intervencao (col 17) e Causas (col 21) a partir da Critica
    if sv and (sv.get('oc_causa') or sv.get('oc_num')):
        ini=str(sv.get('oc_ini') or '')
        try:
            dt=datetime.datetime.strptime(ini,'%Y-%m-%d %H:%M'); ini=dt.strftime('%d/%m/%Y as %H:%M')
        except: pass
        try:
            hh=float(sv.get('oc_dur_h') or 0); dur=f' ~{int(hh):02d}:{int(round((hh-int(hh))*60)):02d}'
        except: dur=''
        if sv.get('oc_num'): put(17,f'Número: {sv["oc_num"]}\nData de abertura: {ini}{dur}')
        cau=str(sv.get('oc_causa') or '').strip(); sub=str(sv.get('oc_sub') or '').strip()
        put(21,f'{cau.capitalize()} — {sub.capitalize()}' if sub else cau.capitalize())
    elif r and V('NUM_INT'):
        put(17,f'Número: {V("NUM_INT")}')
COL={17:'Intervenção',18:'NUMERO_OS',19:'ORIGEM OS',20:'DEFEITO OS',21:'Causas',22:'Anotações OS',
     23:'Dados Trafo retirado',24:'Dados Trafo Instalado',25:'Aterramento',26:'infos gerais',27:'NUM_OBRA'}
print('CÉLULAS PREENCHIDAS:')
for c in sorted(preenchidos): print(f'  col {c:2d} {COL.get(c,""):24s} {preenchidos[c]:2d}')
wb.save('analise_trafos_rurais_1125_e_150_PREENCHIDA.xlsx')
print('\nsalvo')
